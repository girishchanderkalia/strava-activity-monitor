import json
from collections import defaultdict
from openpyxl import load_workbook

WORKBOOK = r"C:\Users\giris\OneDrive\fitness\walking challenge\Strava_Challenge_Production.xlsx"


def update_excel():

    with open("data/activities.json") as f:
        activities = json.load(f)

    wb = load_workbook(WORKBOOK)

    activities_sheet = wb["Activities"]
    leaderboard_sheet = wb["Leaderboard"]

    #
    # Clear Activities tab
    #
    activities_sheet.delete_rows(
        2,
        activities_sheet.max_row
    )

    #
    # Populate Activities tab
    #
    row = 2

    for activity in activities:

        activities_sheet.cell(
            row=row,
            column=1,
            value=activity["date"]
        )

        activities_sheet.cell(
            row=row,
            column=2,
            value=activity["athlete_id"]
        )

        activities_sheet.cell(
            row=row,
            column=3,
            value=activity["name"]
        )

        activities_sheet.cell(
            row=row,
            column=4,
            value=activity["activity_id"]
        )

        activities_sheet.cell(
            row=row,
            column=5,
            value="Walk"
        )

        activities_sheet.cell(
            row=row,
            column=6,
            value=activity["distance_km"]
        )

        row += 1

    #
    # Build leaderboard
    #

    totals = defaultdict(float)

    for activity in activities:

        totals[
            activity["name"]
        ] += activity["distance_km"]

    leaderboard_sheet.delete_rows(
        2,
        leaderboard_sheet.max_row
    )

    leaderboard = sorted(
        totals.items(),
        key=lambda x: x[1],
        reverse=True
    )

    row = 2

    for name, km in leaderboard:

        leaderboard_sheet.cell(
            row=row,
            column=1,
            value=name
        )

        leaderboard_sheet.cell(
            row=row,
            column=2,
            value=round(km, 2)
        )

        row += 1

    wb.save(WORKBOOK)

    print(
        "Excel updated successfully"
    )


if __name__ == "__main__":
    update_excel()